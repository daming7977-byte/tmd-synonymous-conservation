from pathlib import Path
import os
from collections import defaultdict
import gzip
import math
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

try:
    import statsmodels.api as sm
except Exception as e:
    raise SystemExit("STATSMODELS_IMPORT_FAILED: %r" % (e,))


# ============================================================
# PATHS
# ============================================================

SOURCE_PARENT = Path(os.environ["P5_DATA_PARENT"])

ROOT = (
    SOURCE_PARENT
    / "paper5_evolution_spatial_bridge"
    / "05_GSE297497_MPT_BRIDGE"
)

HANDOFF = (
    ROOT
    / "03_TMD_STRUCTURAL_CORRESPONDENCE"
    / "P5_V2_STAGE0E2_SCORED_INTERNAL_TMD_HANDOFF.tsv"
)

EVO = (
    SOURCE_PARENT
    / "paper5_evolution_spatial_bridge"
    / "04_STAGE1_EVOLUTION"
    / "03_FULL_EVOLUTION_DATA"
    / "P5_STAGE1A_FULL_TMD_EVOLUTION_SCORE.tsv"
)

SOURCE_XLSX = (
    ROOT
    / "01_METADATA_SUPP_AUDIT"
    / "PUBLICATION_SOURCE_DATA"
    / "41594_2025_1691_MOESM3_ESM.xlsx"
)

REFGENE = (
    ROOT
    / "09_FORMAL_REFERENCE"
    / "P5_V2_CURRENT_UCSC_hg38_refGene.txt.gz"
)

RUNROOT = (
    ROOT
    / "12_FORMAL_MPT_FULL"
    / "runs"
)

SPEC = (
    ROOT
    / "07_PREREVEAL_FREEZE"
    / "P5_V2_PREREVEAL_MPT_MODEL_SPEC.yaml"
)

OUT = Path.cwd()

OCCDIR = OUT / "occupancy_npz"
OCCDIR.mkdir(exist_ok=True)


# ============================================================
# FROZEN CONSTANTS
# ============================================================

EXPECTED_SPEC_SHA = (
    "181ed0bef6bc81e4750f421c1e0b3f79"
    "a7bfd5fc55432d60664a191d7854b912"
)

PROFILES = {
    "TMCO1_rep1": ("SRR33621384", "SRR33621383"),
    "TMCO1_rep2": ("SRR33621382", "SRR33621381"),
    "CCDC47_rep1": ("SRR33621380", "SRR33621379"),
    "CCDC47_rep2": ("SRR33621378", "SRR33621377"),
    "Nicalin_rep1": ("SRR33621376", "SRR33621375"),
    "Nicalin_rep2": ("SRR33621374", "SRR33621373"),
}

FACTORS = {
    "TMCO1": ["TMCO1_rep1", "TMCO1_rep2"],
    "CCDC47": ["CCDC47_rep1", "CCDC47_rep2"],
    "Nicalin": ["Nicalin_rep1", "Nicalin_rep2"],
}

ALL_RUNS = sorted(
    {
        r
        for pair in PROFILES.values()
        for r in pair
    }
)

CANONICAL = {
    *(f"chr{i}" for i in range(1, 23)),
    "chrX",
    "chrY",
    "chrM",
}


# ============================================================
# UTILITIES
# ============================================================

def sha256_file(path):
    import hashlib

    h = hashlib.sha256()

    with open(path, "rb") as f:
        while True:
            b = f.read(1024 * 1024)
            if not b:
                break
            h.update(b)

    return h.hexdigest()


def qc_value(run, key):
    p = (
        RUNROOT
        / run
        / "final"
        / f"{run}.QC.txt"
    )

    d = {}

    for line in p.read_text().splitlines():
        if " = " in line:
            k, v = line.split(" = ", 1)
            d[k.strip()] = v.strip()

    return d[key]


def zscore_outcome_independent(x):
    x = pd.to_numeric(
        x,
        errors="coerce",
    ).astype(float)

    mu = np.nanmean(x)
    sd = np.nanstd(x, ddof=0)

    if not np.isfinite(sd) or sd == 0:
        raise RuntimeError(
            "ZERO_OR_INVALID_SD"
        )

    return (x - mu) / sd


def holm_adjust(pvals):
    pvals = np.asarray(
        pvals,
        dtype=float,
    )

    m = len(pvals)
    order = np.argsort(pvals)

    out = np.empty(m, dtype=float)

    running = 0.0

    for rank, idx in enumerate(order):
        val = (
            pvals[idx]
            * (m - rank)
        )

        val = min(
            1.0,
            max(
                running,
                val,
            ),
        )

        running = val
        out[idx] = val

    return out


# ============================================================
# LOCKED SPEC CHECK
# ============================================================

actual_spec_sha = sha256_file(
    SPEC
)

print(
    "PREREVEAL_SPEC_SHA256 =",
    actual_spec_sha
)

if actual_spec_sha != EXPECTED_SPEC_SHA:
    raise RuntimeError(
        "PREREVEAL_SPEC_SHA_MISMATCH"
    )

print(
    "PREREVEAL_SPEC_IDENTITY_GATE = GREEN"
)


# ============================================================
# REQUIRE ALL 12 FORMAL RUNS
# ============================================================

for run in ALL_RUNS:

    done = RUNROOT / run / ".DONE"

    asite = (
        RUNROOT
        / run
        / "final"
        / f"{run}.asite_plus15.tsv"
    )

    qc = (
        RUNROOT
        / run
        / "final"
        / f"{run}.QC.txt"
    )

    if not done.exists():
        raise RuntimeError(
            f"{run}: DONE marker absent"
        )

    if not asite.is_file():
        raise RuntimeError(
            f"{run}: A-site table absent"
        )

    if not qc.is_file():
        raise RuntimeError(
            f"{run}: QC absent"
        )

print(
    "FORMAL_RUNS_PRESENT = 12"
)


# ============================================================
# LOAD TMD HANDOFF + EVOLUTION SCORE
# ============================================================

handoff = pd.read_csv(
    HANDOFF,
    sep="\t",
)

if len(handoff) != 5538:
    raise RuntimeError(
        f"HANDOFF_ROWS={len(handoff)} != 5538"
    )

if not (
    handoff["mapping_status"]
    .astype(str)
    .eq("PASS")
    .all()
):
    raise RuntimeError(
        "NONPASS_TMD_IN_HANDOFF"
    )

handoff["author_tx"] = (
    handoff["author_transcript_names"]
    .astype(str)
    .str.strip()
)

targets = sorted(
    handoff["author_tx"].unique()
)

print(
    "PRELOCKED_TMD =",
    len(handoff)
)

print(
    "TARGET_AUTHOR_TRANSCRIPTS =",
    len(targets)
)

if len(targets) != 1400:
    raise RuntimeError(
        f"TARGET_TX={len(targets)} != 1400"
    )


evo = pd.read_csv(
    EVO,
    sep="\t",
)

evo = evo.rename(
    columns={
        "tmd_index":
            "p5_tmd_index"
    }
)

merge_keys = [
    "human_ensembl_gene",
    "human_ensembl_transcript",
    "p5_tmd_index",
]

keep_evo = merge_keys + [
    "POST30_90_4D_EXCESS_CONSERVATION_Z_MEDIAN",
    "median_window_aa_identity",
    "human_window_GC3",
    "human_window_CpG_density",
]

m = handoff.merge(
    evo[keep_evo],
    on=merge_keys,
    how="left",
    validate="one_to_one",
)

score_col = (
    "POST30_90_4D_EXCESS_CONSERVATION_Z_MEDIAN"
)

if m[score_col].isna().any():
    raise RuntimeError(
        "MISSING_FROZEN_EVOLUTION_SCORE"
    )


# ============================================================
# LOAD DEEPTMHMM TARGET ROWS
# ============================================================

print(
    "LOAD_DEEPTMHMM = START"
)

wb = load_workbook(
    SOURCE_XLSX,
    read_only=True,
    data_only=True,
)

ws = wb["DeepTMHMM topology"]

source_rows = {}

target_set = set(targets)

for row in ws.iter_rows(
    min_row=2,
    values_only=True,
):

    tx = row[1]

    if tx is None:
        continue

    tx = str(tx).strip()

    if tx not in target_set:
        continue

    sequence = (
        str(row[3]).strip()
        if row[3] is not None
        else ""
    )

    segments = []

    for j in range(
        10,
        len(row),
        5,
    ):

        if j + 4 >= len(row):
            break

        label = row[j]
        start = row[j + 2]
        end = row[j + 3]
        length = row[j + 4]

        if (
            label is None
            or start is None
            or end is None
        ):
            continue

        segments.append(
            {
                "label":
                    str(label).strip(),
                "start":
                    int(start),
                "end":
                    int(end),
                "length":
                    (
                        int(length)
                        if length is not None
                        else int(end) - int(start) + 1
                    ),
            }
        )

    record = {
        "gene":
            str(row[0]).strip(),
        "sequence":
            sequence,
        "protein_class":
            str(row[4]).strip(),
        "segments":
            segments,
    }

    if tx in source_rows:
        old = source_rows[tx]

        if (
            old["sequence"]
            != record["sequence"]
        ):
            raise RuntimeError(
                f"DEEPTMHMM_DUPLICATE_CONFLICT {tx}"
            )

    else:
        source_rows[tx] = record


if set(source_rows) != target_set:

    missing = sorted(
        target_set
        - set(source_rows)
    )

    raise RuntimeError(
        "DEEPTMHMM_TARGET_MISSING: "
        + ",".join(missing[:20])
    )

print(
    "DEEPTMHMM_TARGET_ROWS =",
    len(source_rows)
)


# ============================================================
# LOAD UNIQUE CANONICAL REFGENE MODEL FOR EACH TARGET
# ============================================================

ref_rows = defaultdict(list)

with gzip.open(
    REFGENE,
    "rt",
) as fh:

    for line in fh:

        f = line.rstrip("\n").split("\t")

        if len(f) < 11:
            continue

        tx = f[1]

        if tx not in target_set:
            continue

        chrom = f[2]

        if chrom not in CANONICAL:
            continue

        ref_rows[tx].append(f)


for tx in targets:

    if len(ref_rows[tx]) != 1:
        raise RuntimeError(
            f"{tx}: canonical rows "
            f"{len(ref_rows[tx])} != 1"
        )

print(
    "TARGET_UNIQUE_CANONICAL_REFGENE =",
    len(targets)
)


# ============================================================
# BUILD TRANSCRIPT COORDINATE MODELS
# ============================================================

models = {}

cds_delta_counts = defaultdict(int)

for tx in targets:

    f = ref_rows[tx][0]

    chrom = f[2]
    strand = f[3]

    tx_start = int(f[4])
    tx_end = int(f[5])

    cds_start = int(f[6])
    cds_end = int(f[7])

    exon_count = int(f[8])

    exon_starts = [
        int(x)
        for x in f[9]
        .rstrip(",")
        .split(",")
        if x
    ]

    exon_ends = [
        int(x)
        for x in f[10]
        .rstrip(",")
        .split(",")
        if x
    ]

    if (
        len(exon_starts) != exon_count
        or len(exon_ends) != exon_count
    ):
        raise RuntimeError(
            f"{tx}: exon count mismatch"
        )

    genomic_exons = list(
        zip(
            exon_starts,
            exon_ends,
        )
    )

    if strand == "+":
        ordered = genomic_exons
    elif strand == "-":
        ordered = list(
            reversed(genomic_exons)
        )
    else:
        raise RuntimeError(
            f"{tx}: invalid strand"
        )

    exon_maps = []

    t_cursor = 0

    cds_t_intervals = []

    for gs, ge in ordered:

        length = ge - gs

        exon_maps.append(
            {
                "gstart": gs,
                "gend": ge,
                "tstart": t_cursor,
            }
        )

        is0 = max(
            gs,
            cds_start,
        )

        ie0 = min(
            ge,
            cds_end,
        )

        if is0 < ie0:

            if strand == "+":

                ts = (
                    t_cursor
                    + (is0 - gs)
                )

                te = (
                    t_cursor
                    + (ie0 - gs)
                )

            else:

                ts = (
                    t_cursor
                    + (ge - ie0)
                )

                te = (
                    t_cursor
                    + (ge - is0)
                )

            cds_t_intervals.append(
                (ts, te)
            )

        t_cursor += length

    if not cds_t_intervals:
        raise RuntimeError(
            f"{tx}: no CDS"
        )

    cds_t_start = min(
        x[0]
        for x in cds_t_intervals
    )

    cds_t_end = max(
        x[1]
        for x in cds_t_intervals
    )

    protein_len = len(
        source_rows[tx]["sequence"]
    )

    cds_len = (
        cds_t_end
        - cds_t_start
    )

    required = (
        protein_len * 3
    )

    delta = (
        cds_len
        - required
    )

    cds_delta_counts[delta] += 1

    if cds_len < required:
        raise RuntimeError(
            f"{tx}: CDS {cds_len} "
            f"< protein_nt {required}"
        )

    models[tx] = {
        "chrom": chrom,
        "strand": strand,
        "tx_len": t_cursor,
        "cds_t_start": cds_t_start,
        "cds_t_end": cds_t_end,
        "protein_len": protein_len,
        "exons": exon_maps,
    }


print(
    "CDS_MINUS_PROTEIN_NT_COUNTS =",
    dict(
        sorted(
            cds_delta_counts.items()
        )
    )
)


# ============================================================
# BUILD OUTCOME-INDEPENDENT COVARIATES
# ============================================================

KD = {
    "I": 4.5,
    "V": 4.2,
    "L": 3.8,
    "F": 2.8,
    "C": 2.5,
    "M": 1.9,
    "A": 1.8,
    "G": -0.4,
    "T": -0.7,
    "S": -0.8,
    "W": -0.9,
    "Y": -1.3,
    "P": -1.6,
    "H": -3.2,
    "E": -3.5,
    "Q": -3.5,
    "D": -3.5,
    "N": -3.5,
    "K": -3.9,
    "R": -4.5,
}


tmd_mean_kd = []
down_len = []
down_lumen = []

for row in m.itertuples(
    index=False
):

    tx = row.author_tx

    source = source_rows[tx]

    sequence = source["sequence"]

    start = int(
        row.author_start_1based
    )

    end = int(
        row.author_end_1based
    )

    tmd_seq = sequence[
        start - 1:
        end
    ]

    vals = [
        KD[a]
        for a in tmd_seq
        if a in KD
    ]

    tmd_mean_kd.append(
        np.mean(vals)
        if vals
        else np.nan
    )

    segments = source["segments"]

    tmd_segment_indices = [
        i
        for i, seg in enumerate(segments)
        if seg["label"] == "Transmembrane"
    ]

    ti = int(
        row.author_tmd_index
    )

    if (
        ti < 1
        or ti > len(tmd_segment_indices)
    ):

        down_len.append(np.nan)
        down_lumen.append(np.nan)
        continue

    seg_idx = (
        tmd_segment_indices[
            ti - 1
        ]
    )

    seg = segments[seg_idx]

    if (
        seg["start"] != start
        or seg["end"] != end
    ):
        raise RuntimeError(
            f"{tx} TMD{ti}: "
            "DeepTMHMM coordinate mismatch"
        )

    downstream = None

    for seg2 in segments[
        seg_idx + 1:
    ]:

        if seg2["label"] != "Transmembrane":
            downstream = seg2
            break

    if downstream is None:

        down_len.append(np.nan)
        down_lumen.append(np.nan)

    else:

        down_len.append(
            float(
                downstream["length"]
            )
        )

        if downstream["label"] == "Outside":
            down_lumen.append(1.0)

        elif downstream["label"] == "Inside":
            down_lumen.append(0.0)

        else:
            down_lumen.append(np.nan)


m["TMD_mean_Kyte_Doolittle_hydropathy"] = (
    tmd_mean_kd
)

m["downstream_nonTMD_segment_length_aa"] = (
    down_len
)

m["downstream_segment_lumenal_indicator"] = (
    down_lumen
)

m[
    "log1p_downstream_nonTMD_segment_length_aa"
] = np.log1p(
    m[
        "downstream_nonTMD_segment_length_aa"
    ]
)


# ============================================================
# PRELOCKED STANDARDIZATION — BEFORE OUTCOME
# ============================================================

m["score_z"] = zscore_outcome_independent(
    m[score_col]
)

continuous_covariates_raw = [
    "median_window_aa_identity",
    "human_window_GC3",
    "human_window_CpG_density",
    "p5_length_aa",
    "TMD_mean_Kyte_Doolittle_hydropathy",
    "p5_tmd_index",
    "log1p_downstream_nonTMD_segment_length_aa",
]

continuous_covariates_z = []

for c in continuous_covariates_raw:

    zc = c + "_z"

    m[zc] = zscore_outcome_independent(
        m[c]
    )

    continuous_covariates_z.append(
        zc
    )


print(
    "OUTCOME_INDEPENDENT_COVARIATES_READY = TRUE"
)


# ============================================================
# MAP A-SITES TO TARGET TRANSCRIPTS AND BUILD POSITIONAL RPM
#
# For each A-site-adjusted read:
#   extend transcript-coordinate position +/-50 nt
#   calculate nucleotide coverage
#   codon occupancy = mean coverage across 3 codon nt
#   divide by total assigned A-site reads / 1e6
# ============================================================

def build_run_occupancy(run):

    npz_path = (
        OCCDIR
        / f"{run}.target_codon_rpm.npz"
    )

    if npz_path.exists():

        z = np.load(
            npz_path,
            allow_pickle=False,
        )

        result = {
            k: z[k]
            for k in z.files
        }

        print(
            run,
            "OCCUPANCY_CACHE = REUSED",
        )

        return result

    path = (
        RUNROOT
        / run
        / "final"
        / f"{run}.asite_plus15.tsv"
    )

    total = int(
        qc_value(
            run,
            "ASITE_ASSIGNED_READS",
        )
    )

    if total <= 0:
        raise RuntimeError(
            f"{run}: invalid denominator"
        )

    print(
        run,
        "READ_ASITE_TABLE = START",
        "N =",
        total,
    )

    df = pd.read_csv(
        path,
        sep="\t",
        usecols=[
            "chrom",
            "strand",
            "asite_0based",
        ],
        dtype={
            "chrom": "string",
            "strand": "string",
            "asite_0based":
                np.int64,
        },
    )

    if len(df) != total:
        raise RuntimeError(
            f"{run}: TSV rows {len(df)} "
            f"!= QC {total}"
        )

    grouped = {}

    for (
        chrom,
        strand,
    ), g in df.groupby(
        ["chrom", "strand"],
        sort=False,
    ):

        arr = g[
            "asite_0based"
        ].to_numpy(
            dtype=np.int64,
            copy=True,
        )

        arr.sort()

        grouped[
            (str(chrom), str(strand))
        ] = arr

    del df

    result = {}

    mapped_events = 0

    for n_tx, tx in enumerate(
        targets,
        start=1,
    ):

        model = models[tx]

        arr = grouped.get(
            (
                model["chrom"],
                model["strand"],
            )
        )

        chunks = []

        if arr is not None:

            for ex in model["exons"]:

                gs = ex["gstart"]
                ge = ex["gend"]
                ts = ex["tstart"]

                lo = np.searchsorted(
                    arr,
                    gs,
                    side="left",
                )

                hi = np.searchsorted(
                    arr,
                    ge,
                    side="left",
                )

                if hi <= lo:
                    continue

                pos = arr[
                    lo:hi
                ]

                if model["strand"] == "+":

                    tc = (
                        ts
                        + (
                            pos
                            - gs
                        )
                    )

                else:

                    tc = (
                        ts
                        + (
                            ge
                            - 1
                            - pos
                        )
                    )

                chunks.append(
                    tc.astype(
                        np.int64,
                        copy=False,
                    )
                )

        if chunks:

            coords = np.concatenate(
                chunks
            )

            mapped_events += len(
                coords
            )

            starts = np.maximum(
                coords - 50,
                0,
            )

            ends = np.minimum(
                coords + 51,
                model["tx_len"],
            )

            diff = np.zeros(
                model["tx_len"] + 1,
                dtype=np.int64,
            )

            np.add.at(
                diff,
                starts,
                1,
            )

            np.add.at(
                diff,
                ends,
                -1,
            )

            cov = np.cumsum(
                diff[:-1]
            )

        else:

            cov = np.zeros(
                model["tx_len"],
                dtype=np.int64,
            )

        cs = model["cds_t_start"]

        ncod = model[
            "protein_len"
        ]

        coding = cov[
            cs:
            cs + 3 * ncod
        ]

        if len(coding) != 3 * ncod:
            raise RuntimeError(
                f"{run} {tx}: coding length error"
            )

        codon_cov = (
            coding
            .reshape(
                ncod,
                3,
            )
            .mean(axis=1)
        )

        rpm = (
            codon_cov
            * 1_000_000.0
            / total
        ).astype(
            np.float32
        )

        result[tx] = rpm

        if n_tx % 200 == 0:
            print(
                run,
                "TARGET_TRANSCRIPTS_DONE =",
                n_tx,
            )

    print(
        run,
        "TARGET_TRANSCRIPT_MAPPING_EVENTS =",
        mapped_events,
    )

    np.savez_compressed(
        npz_path,
        **result,
    )

    print(
        run,
        "OCCUPANCY_CACHE = WRITTEN",
    )

    return result


occupancy = {}

for run in ALL_RUNS:

    occupancy[run] = (
        build_run_occupancy(
            run
        )
    )


print(
    "ALL_12_TARGET_OCCUPANCIES_READY = TRUE"
)


# ============================================================
# METHODS TERMINOLOGY / B CONSTANT NOTE
# ============================================================

note = """P5-V2 POSITIONAL ENRICHMENT IMPLEMENTATION NOTE

Publication Methods:
- A-site adjusted reads extended +/-50 nt.
- positional codon occupancy described as reads per million.
- the later formula paragraph refers to Input_i,j / A_j using TPM terminology.

Formal P5-v2 implementation:
- positional IP_i,j and Input_i,j use codon RPM derived from extended reads.
- A_j is mean codon RPM across the protein-coding region of the paired input.

For the current hypothesis-test reveal, the profile-wide term
log2(B_Input/B_IP) is omitted because it is constant across every TMD
within a profile. Therefore it changes only the regression intercept.
It has exactly zero effect on:
- population score beta / SE / p value,
- adjusted score beta / SE / p value,
- gene-demeaned within-protein beta / SE / p value,
- factor-specific or replicate score slopes.

Full B-normalized enrichment values can be reconstructed later for
reporting without changing any inferential result.

MPT_OUTCOME_REVEALED = TRUE only after outcome table below is created.
"""

(
    OUT
    / "P5_V2_POSITIONAL_ENRICHMENT_IMPLEMENTATION_NOTE.txt"
).write_text(note)


# ============================================================
# FIRST OUTCOME CONSTRUCTION
# ============================================================

profile_cols = []

negative_profile_cols = []

for pname, (
    input_run,
    ip_run,
) in PROFILES.items():

    col = (
        pname
        + "_TMD_score"
    )

    negcol = (
        pname
        + "_NEG_score"
    )

    profile_cols.append(col)
    negative_profile_cols.append(
        negcol
    )

    vals = []
    negvals = []

    for row in m.itertuples(
        index=False
    ):

        tx = row.author_tx

        inp = occupancy[
            input_run
        ][tx]

        ip = occupancy[
            ip_run
        ][tx]

        A = float(
            np.mean(inp)
        )

        start = int(
            row.author_start_1based
        )

        end = int(
            row.author_end_1based
        )

        protein_len = len(inp)

        # ----------------------------------------------------
        # Primary +30...+90 from TMD C-end
        #
        # Codon number end+30 -> zero-based end+29.
        # Codon number end+90 -> stop-exclusive end+90.
        # ----------------------------------------------------

        if (
            A > 0
            and end + 90
                <= protein_len
        ):

            s = end + 29
            e = end + 90

            inp_win = inp[s:e]
            ip_win = ip[s:e]

            if (
                len(inp_win) == 61
                and len(ip_win) == 61
            ):

                loge = np.log2(
                    (
                        ip_win
                        + A
                    )
                    /
                    (
                        inp_win
                        + A
                    )
                )

                vals.append(
                    float(
                        np.mean(loge)
                    )
                )

            else:
                vals.append(np.nan)

        else:
            vals.append(np.nan)

        # ----------------------------------------------------
        # Negative control:
        # TMD N-start -90 ... -30
        # ----------------------------------------------------

        if (
            A > 0
            and start - 90 >= 1
        ):

            s = start - 91
            e = start - 30

            inp_win = inp[s:e]
            ip_win = ip[s:e]

            if (
                len(inp_win) == 61
                and len(ip_win) == 61
            ):

                loge = np.log2(
                    (
                        ip_win
                        + A
                    )
                    /
                    (
                        inp_win
                        + A
                    )
                )

                negvals.append(
                    float(
                        np.mean(loge)
                    )
                )

            else:
                negvals.append(
                    np.nan
                )

        else:
            negvals.append(
                np.nan
            )

    m[col] = vals
    m[negcol] = negvals


# Require all six profiles for primary.

m["primary_profiles_valid"] = (
    m[
        profile_cols
    ]
    .notna()
    .sum(axis=1)
)

m["primary_MPT_TMD_score"] = np.where(
    m["primary_profiles_valid"] == 6,
    m[profile_cols].mean(axis=1),
    np.nan,
)


m["negative_profiles_valid"] = (
    m[
        negative_profile_cols
    ]
    .notna()
    .sum(axis=1)
)

m["temporal_negative_score"] = np.where(
    m["negative_profiles_valid"] == 6,
    m[
        negative_profile_cols
    ].mean(axis=1),
    np.nan,
)


# Factor-specific two-replicate means.

for factor, cols in FACTORS.items():

    score_cols = [
        x + "_TMD_score"
        for x in cols
    ]

    m[
        factor
        + "_factor_score"
    ] = np.where(
        m[
            score_cols
        ]
        .notna()
        .sum(axis=1)
        == 2,
        m[
            score_cols
        ].mean(axis=1),
        np.nan,
    )


outcome_path = (
    OUT
    / "P5_V2_TMD_OUTCOME_TABLE.tsv"
)

m.to_csv(
    outcome_path,
    sep="\t",
    index=False,
)

print(
    "MPT_OUTCOME_REVEALED = TRUE"
)

print(
    "PRIMARY_VALID_TMD =",
    int(
        m[
            "primary_MPT_TMD_score"
        ]
        .notna()
        .sum()
    )
)

print(
    "TEMPORAL_NEGATIVE_VALID_TMD =",
    int(
        m[
            "temporal_negative_score"
        ]
        .notna()
        .sum()
    )
)


# ============================================================
# CLUSTER-ROBUST OLS
# ============================================================

def cluster_ols(
    df,
    ycol,
    xcols,
    score_name="score_z",
):

    cols = (
        [ycol]
        + list(xcols)
        + [
            "human_ensembl_gene"
        ]
    )

    d = (
        df[cols]
        .replace(
            [np.inf, -np.inf],
            np.nan,
        )
        .dropna()
        .copy()
    )

    X = sm.add_constant(
        d[xcols].astype(float),
        has_constant="add",
    )

    y = d[ycol].astype(float)

    fit = sm.OLS(
        y,
        X,
    ).fit(
        cov_type="cluster",
        cov_kwds={
            "groups":
                d[
                    "human_ensembl_gene"
                ],
            "use_correction":
                True,
        },
        use_t=True,
    )

    ci = fit.conf_int().loc[
        score_name
    ]

    return {
        "N":
            len(d),
        "clusters":
            d[
                "human_ensembl_gene"
            ].nunique(),
        "beta":
            float(
                fit.params[
                    score_name
                ]
            ),
        "se":
            float(
                fit.bse[
                    score_name
                ]
            ),
        "ci_low":
            float(ci.iloc[0]),
        "ci_high":
            float(ci.iloc[1]),
        "p":
            float(
                fit.pvalues[
                    score_name
                ]
            ),
    }


def within_gene_model(df):

    binary = (
        "downstream_segment_lumenal_indicator"
    )

    xcols = (
        ["score_z"]
        + continuous_covariates_z
        + [binary]
    )

    cols = (
        [
            "primary_MPT_TMD_score",
            "human_ensembl_gene",
        ]
        + xcols
    )

    d = (
        df[cols]
        .replace(
            [np.inf, -np.inf],
            np.nan,
        )
        .dropna()
        .copy()
    )

    counts = (
        d.groupby(
            "human_ensembl_gene"
        )
        .size()
    )

    valid_genes = set(
        counts[
            counts >= 2
        ].index
    )

    d = d[
        d[
            "human_ensembl_gene"
        ]
        .isin(
            valid_genes
        )
    ].copy()

    demean_cols = (
        [
            "primary_MPT_TMD_score"
        ]
        + xcols
    )

    for c in demean_cols:

        d[c + "_DM"] = (
            d[c]
            - d.groupby(
                "human_ensembl_gene"
            )[c].transform(
                "mean"
            )
        )

    ycol = (
        "primary_MPT_TMD_score_DM"
    )

    xdm = [
        c + "_DM"
        for c in xcols
    ]

    X = d[xdm].astype(float)
    y = d[ycol].astype(float)

    # Drop nuisance columns with no within variation.
    keep = []

    for c in X.columns:

        if np.nanstd(
            X[c].to_numpy()
        ) > 0:
            keep.append(c)

    X = X[keep]

    if "score_z_DM" not in X.columns:
        raise RuntimeError(
            "WITHIN_SCORE_NO_VARIATION"
        )

    fit = sm.OLS(
        y,
        X,
    ).fit(
        cov_type="cluster",
        cov_kwds={
            "groups":
                d[
                    "human_ensembl_gene"
                ],
            "use_correction":
                True,
        },
        use_t=True,
    )

    ci = fit.conf_int().loc[
        "score_z_DM"
    ]

    return {
        "N":
            len(d),
        "clusters":
            d[
                "human_ensembl_gene"
            ].nunique(),
        "beta":
            float(
                fit.params[
                    "score_z_DM"
                ]
            ),
        "se":
            float(
                fit.bse[
                    "score_z_DM"
                ]
            ),
        "ci_low":
            float(ci.iloc[0]),
        "ci_high":
            float(ci.iloc[1]),
        "p":
            float(
                fit.pvalues[
                    "score_z_DM"
                ]
            ),
    }


# ============================================================
# PRIMARY
# ============================================================

primary = cluster_ols(
    m,
    "primary_MPT_TMD_score",
    ["score_z"],
)


# ============================================================
# ADJUSTED
# ============================================================

adjusted_x = (
    ["score_z"]
    + continuous_covariates_z
    + [
        "downstream_segment_lumenal_indicator"
    ]
)

adjusted = cluster_ols(
    m,
    "primary_MPT_TMD_score",
    adjusted_x,
)


# ============================================================
# WITHIN-PROTEIN
# ============================================================

within = within_gene_model(
    m
)


# ============================================================
# FACTOR-SPECIFIC SECONDARY
# Use the primary-valid universe for direct comparability.
# ============================================================

primary_universe = m[
    m[
        "primary_MPT_TMD_score"
    ].notna()
].copy()


factor_results = {}

factor_p = []

for factor in [
    "TMCO1",
    "CCDC47",
    "Nicalin",
]:

    ycol = (
        factor
        + "_factor_score"
    )

    r = cluster_ols(
        primary_universe,
        ycol,
        ["score_z"],
    )

    factor_results[
        factor
    ] = r

    factor_p.append(
        r["p"]
    )


factor_holm = holm_adjust(
    factor_p
)

for (
    factor,
    ph,
) in zip(
    [
        "TMCO1",
        "CCDC47",
        "Nicalin",
    ],
    factor_holm,
):

    factor_results[
        factor
    ][
        "holm_p"
    ] = float(ph)


# ============================================================
# REPLICATE SENSITIVITY
# ============================================================

replicate_results = {}

for pname in PROFILES:

    ycol = (
        pname
        + "_TMD_score"
    )

    replicate_results[
        pname
    ] = cluster_ols(
        primary_universe,
        ycol,
        ["score_z"],
    )


# ============================================================
# TEMPORAL NEGATIVE CONTROL
# Compare where BOTH primary and negative are valid.
# ============================================================

neg_universe = m[
    m[
        "primary_MPT_TMD_score"
    ].notna()
    &
    m[
        "temporal_negative_score"
    ].notna()
].copy()

negative = cluster_ols(
    neg_universe,
    "temporal_negative_score",
    ["score_z"],
)


# ============================================================
# FORMAL ADJUDICATION
# ============================================================

core_p = [
    primary["p"],
    adjusted["p"],
    within["p"],
]

if all(
    p >= 0.05
    for p in core_p
):

    adjudication = (
        "STOP_PRIMARY_P5_V2"
    )

elif all(
    p < 0.05
    for p in core_p
):

    adjudication = (
        "STRONG_BRIDGE_SUPPORT"
    )

else:

    adjudication = (
        "INTERMEDIATE_BRIDGE_RESULT"
    )


primary_sign = np.sign(
    primary["beta"]
)

factor_same_direction = all(
    np.sign(
        factor_results[f][
            "beta"
        ]
    )
    == primary_sign
    for f in factor_results
)

rep_same_direction_count = sum(
    np.sign(
        replicate_results[r][
            "beta"
        ]
    )
    == primary_sign
    for r in replicate_results
)

negative_weaker = (
    abs(
        negative["beta"]
    )
    <
    abs(
        primary["beta"]
    )
)


# ============================================================
# WRITE RESULT TABLES
# ============================================================

def result_row(name, r):
    return {
        "analysis":
            name,
        **r,
    }


core_df = pd.DataFrame(
    [
        result_row(
            "PRIMARY",
            primary,
        ),
        result_row(
            "ADJUSTED",
            adjusted,
        ),
        result_row(
            "WITHIN_PROTEIN",
            within,
        ),
        result_row(
            "TEMPORAL_NEGATIVE",
            negative,
        ),
    ]
)

core_df.to_csv(
    OUT
    / "P5_V2_REVEAL_CORE_RESULTS.tsv",
    sep="\t",
    index=False,
)


factor_df = pd.DataFrame(
    [
        {
            "factor": f,
            **factor_results[f],
        }
        for f in [
            "TMCO1",
            "CCDC47",
            "Nicalin",
        ]
    ]
)

factor_df.to_csv(
    OUT
    / "P5_V2_REVEAL_FACTOR_RESULTS.tsv",
    sep="\t",
    index=False,
)


rep_df = pd.DataFrame(
    [
        {
            "profile": p,
            **replicate_results[p],
        }
        for p in PROFILES
    ]
)

rep_df.to_csv(
    OUT
    / "P5_V2_REVEAL_REPLICATE_RESULTS.tsv",
    sep="\t",
    index=False,
)


# ============================================================
# PRINT FINAL REVEAL
# ============================================================

def fmt(r):

    return (
        "N={N} genes={clusters} "
        "beta={beta:.6g} "
        "SE={se:.6g} "
        "95%CI=[{ci_low:.6g},{ci_high:.6g}] "
        "p={p:.6g}"
    ).format(**r)


lines = []

lines.append(
    "============================================================"
)

lines.append(
    "P5-V2 FORMAL REVEAL"
)

lines.append(
    "============================================================"
)

lines.append("")

lines.append(
    "PRIMARY: "
    + fmt(primary)
)

lines.append(
    "ADJUSTED: "
    + fmt(adjusted)
)

lines.append(
    "WITHIN_PROTEIN: "
    + fmt(within)
)

lines.append("")

lines.append(
    "TEMPORAL_NEGATIVE: "
    + fmt(negative)
)

lines.append("")

lines.append(
    "FACTOR_SPECIFIC:"
)

for f in [
    "TMCO1",
    "CCDC47",
    "Nicalin",
]:

    r = factor_results[f]

    lines.append(
        "  %s: %s Holm_p=%.6g"
        % (
            f,
            fmt(r),
            r["holm_p"],
        )
    )

lines.append("")

lines.append(
    "REPLICATE_SENSITIVITY:"
)

for p in PROFILES:

    lines.append(
        "  %s: %s"
        % (
            p,
            fmt(
                replicate_results[p]
            ),
        )
    )

lines.append("")

lines.append(
    "FACTOR_ALL_SAME_DIRECTION_AS_PRIMARY = "
    + str(
        factor_same_direction
    ).upper()
)

lines.append(
    "REPLICATES_SAME_DIRECTION_AS_PRIMARY = "
    + str(
        rep_same_direction_count
    )
    + "/6"
)

lines.append(
    "TEMPORAL_NEGATIVE_ABS_BETA_WEAKER_THAN_PRIMARY = "
    + str(
        negative_weaker
    ).upper()
)

lines.append("")

lines.append(
    "FORMAL_P5_V2_ADJUDICATION = "
    + adjudication
)

lines.append("")

if adjudication == "STOP_PRIMARY_P5_V2":

    lines.append(
        "ACTION = STOP; NO WINDOW SCAN; "
        "NO PREDICTOR REDEFINITION; "
        "NO OST-A RESCUE."
    )

elif adjudication == "STRONG_BRIDGE_SUPPORT":

    lines.append(
        "ACTION = CONTINUE P5-V2 AS SUPPORTED BRIDGE; "
        "EVALUATE STORY STRENGTH / SPECIFICITY NEXT."
    )

else:

    lines.append(
        "ACTION = PRESERVE AS INTERMEDIATE; "
        "DO NOT RETUNE PRIMARY WINDOW OR PREDICTOR."
    )

lines.append("")

lines.append(
    "CAUSAL_CLAIM_AUTHORIZED = FALSE"
)

text = "\n".join(lines)

print()
print(text)

(
    OUT
    / "P5_V2_FORMAL_REVEAL.txt"
).write_text(
    text + "\n"
)

print()
print(
    "REVEAL_OUTPUTS_WRITTEN = TRUE"
)
