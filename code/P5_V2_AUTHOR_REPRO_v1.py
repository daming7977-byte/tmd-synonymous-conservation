#!/usr/bin/env python3
"""Validation 1 only. Explicit inputs; no evolutionary/outcome tables read.
Run with bundled Python; dependencies are confined beside this file.
Primary method is fixed in METHOD_ADJUDICATION before counting/comparison.
"""
from pathlib import Path
import sys, os, re, gzip, json, csv, hashlib, subprocess, datetime, collections
from concurrent.futures import ThreadPoolExecutor, as_completed
OUT=Path(__file__).resolve().parent
sys.path.insert(0,str(OUT/'dependencies'))
os.environ['PYTHONPATH']=str(OUT/'dependencies')
os.environ['PYTHONDONTWRITEBYTECODE']='1'
import numpy as np
import openpyxl, pysam, HTSeq
from scipy import stats
ROOT=OUT.parent.parent
PREFIX='P5_V2_AUTHOR_REPRO_'
GTF=ROOT/'09_FORMAL_REFERENCE/P5_V2_FORMAL_CURRENT_refGene.gtf'
REF=ROOT/'09_FORMAL_REFERENCE/P5_V2_CURRENT_UCSC_hg38_refGene.txt.gz'
REP=ROOT/'06_AUTHOR_PIPELINE_RECONCILIATION/P5_V2_STAGE0H_AUTHOR_REPRESENTATIVE_TRANSCRIPTS.txt'
BOOK=ROOT/'01_METADATA_SUPP_AUDIT/PUBLICATION_SOURCE_DATA/41594_2025_1691_MOESM3_ESM.xlsx'
REPO=ROOT/'06_AUTHOR_PIPELINE_RECONCILIATION/RNaseFootprinting/README.md'
CHARTER=OUT.parent/'P5_V2_PUBLICATION_VALIDATION_CHARTER.md'
PROFILES=[('MPT','TMCO1-R1',84,83),('MPT','TMCO1-R2',82,81),('MPT','CCDC47-R1',80,79),('MPT','CCDC47-R2',78,77),('MPT','Nicalin-R1',76,75),('MPT','Nicalin-R2',74,73),('OST-A','OST48-R1',90,89),('OST-A','OST48-R2',88,87),('OST-A','RPN2-R1',86,85)]

def log(s): print(datetime.datetime.now(datetime.timezone.utc).isoformat(),s,flush=True)
def sha(p):
 h=hashlib.sha256()
 with open(p,'rb') as f:
  for b in iter(lambda:f.read(8*1024*1024),b''):h.update(b)
 return h.hexdigest()
def write_tsv(p,rows,fields=None):
 rows=list(rows)
 with (gzip.open(p,'wt',newline='') if str(p).endswith('.gz') else open(p,'w',newline='')) as f:
  w=csv.DictWriter(f,fieldnames=fields or list(rows[0]),delimiter='\t');w.writeheader();w.writerows(rows)
def bam_path(n):
 s=f'SRR336213{n}';area='12_FORMAL_MPT_FULL' if n<=84 else '14_OSTA_SPECIFICITY'
 return ROOT/area/'runs'/s/'final'/f'{s}.selected_lengths.bam'
def source_audit():
 w=openpyxl.load_workbook(BOOK,read_only=True,data_only=True)
 topology=collections.defaultdict(set);tx_by_gene=collections.defaultdict(set)
 for gene,tx,length,seq,cls in w['DeepTMHMM topology'].iter_rows(min_row=2,max_col=5,values_only=True):
  if gene:
   topology[gene].add(cls);tx_by_gene[gene].add(tx)
 assert all(len(t)==1 for t in tx_by_gene.values())
 reps=set(REP.read_text().split());assert reps==set.union(*tx_by_gene.values())
 sources={};audit={};background={g for g,c in topology.items() if c=={'CYTONUC'}}
 for sheet in ['MPT','OST-A']:
  it=w[sheet].iter_rows(values_only=True);headers=next(it)
  # Only authorized count, class and enrichment values are retained; no p-values.
  keep=[i for i,h in enumerate(headers) if h and (h=='Gene' or h=='DeepTMHMM protein class' or h.startswith('Normalized transcript count ') or '(log2)' in h)]
  rows=[{headers[i]:r[i] for i in keep} for r in it if r[0]]
  assert len(rows)==len({r['Gene'] for r in rows});sources[sheet]={r['Gene']:r for r in rows}
  audit[sheet]={'genes':len(rows),'CYTONUC_rows':sum(r['DeepTMHMM protein class']=='CYTONUC' for r in rows),'profiles':{}}
  for _,pr,inp,ip in [p for p in PROFILES if p[0]==sheet]:
   factor,rep=pr.split('-');ic='Normalized transcript count '+factor+'-IP-'+rep;nc='Normalized transcript count '+factor+'-Input-'+rep
   ec=next(k for k in rows[0] if k.startswith(pr+' ') and '(log2)' in k)
   assert all(isinstance(r[ic],(int,float)) and r[ic]>0 and isinstance(r[nc],(int,float)) and r[nc]>0 for r in rows)
   residual=max(abs(np.log2(r[ic]/r[nc])-r[ec]) for r in rows)
   audit[sheet]['profiles'][pr]={'enrichment_column':ec,'IP_column':ic,'Input_column':nc,'max_log_ratio_residual':residual,'published_CYTONUC_median':float(np.median([r[ec] for r in rows if r['DeepTMHMM protein class']=='CYTONUC']))}
  cols=[audit[sheet]['profiles'][p[1]]['enrichment_column'] for p in PROFILES if p[0]==sheet]
  av=sheet+' average enrichment (log2)'
  audit[sheet]['combined_max_arithmetic_residual']=max(abs(np.mean([r[c] for c in cols])-r[av]) for r in rows)
  assert audit[sheet]['combined_max_arithmetic_residual']<1e-10
 audit['topology']={'unique_genes':len(topology),'unique_transcripts':len(reps),'CYTONUC_unambiguous':len(background),'conflicting_classes':{g:sorted(c) for g,c in topology.items() if len(c)>1}}
 w.close();return sources,audit,background,tx_by_gene,reps

def annotation(tx_by_gene,reps):
 source_ids=set();source_symbols=collections.defaultdict(set)
 with gzip.open(REF,'rt') as f:
  for l in f:
   a=l.rstrip().split('\t');source_ids.add(a[1]);source_symbols[a[1]].add(a[12])
 tx2gene={next(iter(t)):g for g,t in tx_by_gene.items()};assert len(tx2gene)==len(tx_by_gene)
 for t,g in tx2gene.items():assert g in source_symbols[t]
 intervals=collections.defaultdict(list);seen=set();cdsrows=0;suffixed=0;gtfout=OUT/'representative_CDS.gtf'
 with open(GTF) as f,open(gtfout,'w') as o:
  for l in f:
   if l.startswith('#'):continue
   a=l.rstrip().split('\t')
   if a[2]!='CDS':continue
   attr=dict(re.findall(r'(\w+) "([^"]+)"',a[8]));t=attr['transcript_id'];base=t
   if t not in source_ids:
    b=re.sub(r'_\d+$','',t)
    if b in source_ids:base=b
   if base not in reps:continue
   g=tx2gene[base];assert attr['gene_id']==g
   o.write(l);seen.add(base);cdsrows+=1;suffixed+=base!=t
   intervals[g,a[0],a[6]].append((int(a[3])-1,int(a[4])))
 lengths=collections.Counter()
 for (g,ch,st),ints in intervals.items():
  end=-1
  for a,b in sorted(ints):
   lengths[g]+=max(0,b-max(a,end));end=max(end,b)
 missing=sorted(reps-seen)
 write_tsv(OUT/'annotation_gene_audit.tsv',[{'gene':g,'transcript':next(iter(tx_by_gene[g])),'CDS_union_bp':lengths.get(g,0)} for g in sorted(tx_by_gene)])
 return gtfout,dict(lengths),{'CDS_rows':cdsrows,'suffixed_CDS_rows':suffixed,'CDS_genes':len(lengths),'representatives_without_CDS':missing,'GTF_gene_id':'gene symbol','GTF_transcript_id':'RefSeq accession with validated placement suffix'}

def count_one(n,gtf):
 run=f'SRR336213{n}';bam=bam_path(n);dest=OUT/(run+'.counts.tsv');err=OUT/(run+'.htseq.log');qcfile=OUT/(run+'.bam_qc.json')
 log('START '+run)
 # Inspect every retained read, independent of any previous pipeline QC.
 q=collections.Counter();lens=collections.Counter();maps=collections.Counter()
 with pysam.AlignmentFile(str(bam),'rb') as b:
  for r in b.fetch(until_eof=True):
   q['records']+=1;lens[r.query_length]+=1;maps[r.mapping_quality]+=1
   if r.is_unmapped:q['unmapped']+=1
   if r.is_paired:q['paired']+=1
   if r.is_secondary:q['secondary']+=1
   if r.is_supplementary:q['supplementary']+=1
   if not r.has_tag('NH') or r.get_tag('NH')!=1:q['not_NH1']+=1
 assert set(lens)<={18,19,20,26,27,28,29},(run,lens)
 assert not any(q[x] for x in ['unmapped','paired','secondary','supplementary','not_NH1']),(run,q)
 qc={'run':run,**dict(q),'lengths':dict(lens),'MAPQ':dict(maps)};qcfile.write_text(json.dumps(qc,indent=2)+'\n')
 cmd=[sys.executable,'-m','HTSeq.scripts.count','--format=bam','--order=pos','--stranded=yes','--type=CDS','--idattr=gene_id','--mode=union','--nonunique=all','--minaqual=10',str(bam),str(gtf)]
 with open(dest,'w') as f,open(err,'w') as e:
  e.write('COMMAND '+json.dumps(cmd)+'\n');e.flush();subprocess.run(cmd,stdout=f,stderr=e,check=True,cwd=OUT)
 counts={}
 for l in dest.read_text().splitlines():
  g,c=l.split('\t');counts[g]=int(c)
 assert sum(v for k,v in counts.items() if not k.startswith('__'))>0
 log('DONE '+run+' records='+str(q['records']))
 return n,counts,qc

def metrics(label,sheet,rows):
 a=np.array([r['author_enrichment'] for r in rows]);b=np.array([r['reconstructed_enrichment'] for r in rows]);ok=np.isfinite(a)&np.isfinite(b);a=a[ok];b=b[ok]
 lr=stats.linregress(a,b);res=b-a
 # This is a discrepancy description, not a data-selection or tuning rule.
 return {'profile':label,'group':sheet,'matched_N':len(a),'Pearson_r':float(stats.pearsonr(a,b).statistic),'Spearman_rho':float(stats.spearmanr(a,b).statistic),'OLS_slope_reconstruction_on_author':float(lr.slope),'OLS_intercept':float(lr.intercept),'median_absolute_difference':float(np.median(abs(res))),'median_signed_difference':float(np.median(res)),'residual_p05':float(np.quantile(res,.05)),'residual_p95':float(np.quantile(res,.95)),'sign_concordance_nonzero':float(np.mean(np.sign(a[(a!=0)&(b!=0)])==np.sign(b[(a!=0)&(b!=0)]))),'systematic_discrepancy':f'slope={lr.slope:.6f}; intercept={lr.intercept:.6f}; median_shift={np.median(res):.6f}'}

def main():
 log('VALIDATION_1_ONLY; HTSeq='+HTSeq.__version__+'; Python='+sys.version.split()[0]);assert HTSeq.__version__=='2.0.3'
 sources,audit,bg,txgene,reps=source_audit();gtf,lengths,ann=annotation(txgene,reps);audit['annotation']=ann
 inputs=[GTF,REF,REP,BOOK,REPO,CHARTER]+[bam_path(n) for n in range(73,91)]+[Path(str(bam_path(n))+'.bai') for n in range(73,91)]
 inputs+= [ROOT/'09_FORMAL_REFERENCE/P5_V2_STAGE9D_FORMAL_GTF_FINAL_DECISION.txt',ROOT/'09_FORMAL_REFERENCE/P5_V2_STAGE9D_GTF_DUPLICATE_ID_ERRATUM.txt']
 hashes={str(p):sha(p) for p in inputs};(OUT/(PREFIX+'INPUT_MANIFEST_SHA256.txt')).write_text(''.join(h+'  '+p+'\n' for p,h in hashes.items()))
 log('Input hashes recorded; annotation='+json.dumps(ann))
 (OUT/'source_annotation_audit.json').write_text(json.dumps(audit,indent=2)+'\n')
 results={};qcs={}
 with ThreadPoolExecutor(max_workers=6) as pool:
  for fut in as_completed([pool.submit(count_one,n,gtf) for n in range(73,91)]):
   n,c,q=fut.result();results[n]=c;qcs[n]=q
 genes=sorted(lengths);lens=np.array([lengths[g] for g in genes],float);idx={g:i for i,g in enumerate(genes)}
 counts={n:np.array([results[n][g] for g in genes],dtype=np.int64) for n in results}
 tpm={n:(c/lens)/(c/lens).sum()*1e6 for n,c in counts.items()}
 assert all(abs(t.sum()-1e6)<1e-6 for t in tpm.values())
 write_tsv(OUT/'CDS_COUNTS_TPM.tsv.gz',({'run':f'SRR336213{n}','gene':g,'CDS_union_bp':lengths[g],'count':int(counts[n][j]),'TPM':float(tpm[n][j])} for n in sorted(counts) for j,g in enumerate(genes)))
 matches=[];profiles=[];enrich={};centers=[];allgene=[]
 for sheet,pr,inp,ip in PROFILES:
  valid=(counts[ip]>0)&(counts[inp]>0);raw=np.full(len(genes),np.nan)
  raw[(counts[ip]==0)&(counts[inp]>0)]=-np.inf;raw[(counts[ip]>0)&(counts[inp]==0)]=np.inf
  raw[valid]=np.log2(tpm[ip][valid]/tpm[inp][valid]);bmask=valid&np.array([g in bg for g in genes]);assert bmask.sum()>0
  center=float(np.median(raw[bmask]));en=raw-center;enrich[pr]=en
  assert abs(np.median(en[bmask]))<1e-12
  # Length and TPM scale cancellation independently verifies centered enrichment arithmetic.
  countlog=np.log2(counts[ip][valid]/counts[inp][valid]);bgcount=np.log2(counts[ip][bmask]/counts[inp][bmask]);assert np.max(abs(en[valid]-(countlog-np.median(bgcount))))<1e-10
  s=sources[sheet];ec=audit[sheet]['profiles'][pr]['enrichment_column']
  rows=[{'group':sheet,'profile':pr,'gene':g,'author_enrichment':s[g][ec],'reconstructed_enrichment':float(en[idx[g]]),'difference':float(en[idx[g]]-s[g][ec]),'Input_count':int(counts[inp][idx[g]]),'IP_count':int(counts[ip][idx[g]]),'Input_TPM':float(tpm[inp][idx[g]]),'IP_TPM':float(tpm[ip][idx[g]]),'raw_log2_IP_Input':float(raw[idx[g]]),'CYTONUC_center':center} for g in s if g in idx and valid[idx[g]]]
  matches.extend(rows);m=metrics(pr,sheet,rows);profiles.append(m)
  centers.append({'profile':pr,'positive_pair_N':int(valid.sum()),'CYTONUC_background_N':int(bmask.sum()),'CYTONUC_median_raw_log2':center,'author_N':len(s),'matched_N':len(rows),'author_missing_annotation':sum(g not in idx for g in s),'author_zero_count_excluded':sum(g in idx and not valid[idx[g]] for g in s),'IP_zero_only':int(((counts[ip]==0)&(counts[inp]>0)).sum()),'Input_zero_only':int(((counts[ip]>0)&(counts[inp]==0)).sum()),'both_zero':int(((counts[ip]==0)&(counts[inp]==0)).sum())})
  for j,g in enumerate(genes):allgene.append({'profile':pr,'gene':g,'raw_log2_IP_Input':raw[j],'CYTONUC_center':center,'centered_enrichment':en[j],'finite_pair':bool(valid[j]),'centering_background':bool(bmask[j])})
  log('PROFILE '+json.dumps(m))
 for sheet in ['MPT','OST-A']:
  names=[p[1] for p in PROFILES if p[0]==sheet];stack=np.array([enrich[n] for n in names]);valid=np.all(np.isfinite(stack),axis=0);av=np.full(len(genes),np.nan);av[valid]=np.mean(stack[:,valid],axis=0);s=sources[sheet];pr=sheet+'-average';col=sheet+' average enrichment (log2)'
  rows=[{'group':sheet,'profile':pr,'gene':g,'author_enrichment':s[g][col],'reconstructed_enrichment':float(av[idx[g]]),'difference':float(av[idx[g]]-s[g][col]),'Input_count':'','IP_count':'','Input_TPM':'','IP_TPM':'','raw_log2_IP_Input':'','CYTONUC_center':''} for g in s if g in idx and valid[idx[g]]]
  matches.extend(rows);profiles.append(metrics(pr,sheet,rows));log('COMBINED '+json.dumps(profiles[-1]))
 write_tsv(OUT/(PREFIX+'PROFILE_RESULTS.tsv'),profiles);write_tsv(OUT/(PREFIX+'MATCHED_VALUES.tsv.gz'),matches)
 write_tsv(OUT/'CENTERING_AND_COVERAGE.tsv',centers);write_tsv(OUT/'ALL_GENE_ENRICHMENT.tsv.gz',allgene)
 for p,h in hashes.items():assert sha(p)==h,'INPUT CHANGED: '+p
 log('All input SHA256 hashes unchanged after execution')
 (OUT/'software_versions.json').write_text(json.dumps({'python':sys.version,'HTSeq':HTSeq.__version__,'pysam':pysam.__version__,'numpy':np.__version__,'samtools':subprocess.check_output(['samtools','--version'],text=True).splitlines()[0]},indent=2)+'\n')
 write_summary(profiles,centers,audit,qcs)
 # Outputs are hashed separately to avoid recursive self-hashing.
 outs=[p for p in OUT.iterdir() if p.is_file() and p.name not in [PREFIX+'RUN.log','OUTPUT_MANIFEST_SHA256.txt']]
 (OUT/'OUTPUT_MANIFEST_SHA256.txt').write_text(''.join(sha(p)+'  '+str(p)+'\n' for p in sorted(outs)))
 log('VALIDATION_1_COMPLETE; no Validation 2/3 undertaken')
 print((OUT/'TERMINAL_SUMMARY.txt').read_text(),end='',flush=True)

def write_summary(profiles,centers,audit,qcs):
 single=profiles[:9];combined=profiles[9:];r=float(np.median([p['Pearson_r'] for p in single]));rho=float(np.median([p['Spearman_rho'] for p in single]));minimum=min(p['Spearman_rho'] for p in single)
 # No pass threshold was supplied. Completion means computable independent reconstructions,
 # not exact author equivalence. Quantitative agreement remains for independent review.
 terminal=f'''AUTHOR_REPRO_IMPLEMENTATION = HTSeq_2.0.3_selected_lengths_representative_CDS_no_pseudocount
MPT_PROFILES_REPRODUCED = 6/6
OSTA_PROFILES_REPRODUCED = 3/3
MEDIAN_PEARSON_R = {r:.6f}
MEDIAN_SPEARMAN_RHO = {rho:.6f}
MIN_PROFILE_SPEARMAN_RHO = {minimum:.6f}
COMBINED_MPT_CORRELATION = Pearson {combined[0]['Pearson_r']:.6f}; Spearman {combined[0]['Spearman_rho']:.6f}
COMBINED_OSTA_CORRELATION = Pearson {combined[1]['Pearson_r']:.6f}; Spearman {combined[1]['Spearman_rho']:.6f}
MATERIAL_SYSTEMATIC_DISCREPANCY = PENDING_QUANTITATIVE_REVIEW
AUTHOR_REPRO_TECHNICAL_STATUS = RECONSTRUCTION_COMPLETE_AUTHOR_EQUIVALENCE_NOT_CERTIFIED
NEXT_ACTION = Independent_Validation_1_review_only
'''
 (OUT/'TERMINAL_SUMMARY.txt').write_text(terminal)
 lines=['# Validation 1 — independent author gene-level reconstruction','',
 'All nine profiles and both source-verified arithmetic averages were independently reconstructed. “Reproduced” in the terminal count means the numerical reconstruction and comparison completed; it is not a declaration of exact equivalence. No acceptance threshold was supplied. Scientific hypothesis testing remains closed.','',
 'Regression uses author enrichment as x and reconstructed enrichment as y. Differences are reconstruction minus author. Correlations use only finite matched values; combined values require all component profiles finite. No correlation p-values were calculated or used.','',
 '| Profile | N | Pearson r | Spearman rho | Slope | Intercept | Median absolute difference |',
 '|---|---:|---:|---:|---:|---:|---:|']
 for p in profiles:lines.append(f"| {p['profile']} | {p['matched_N']} | {p['Pearson_r']:.6f} | {p['Spearman_rho']:.6f} | {p['OLS_slope_reconstruction_on_author']:.6f} | {p['OLS_intercept']:.6f} | {p['median_absolute_difference']:.6f} |")
 lines+=['','## Interpretation and limitations','',
 'Read selection, representative-CDS restriction, HTSeq options, zero handling and background definition were fixed before reconstructed enrichment comparisons. No evolutionary score, P5 association table, alternative TMD window or A-site table was read. BAMs were read directly for gene counting; all output stayed in this directory.','',
 'The historical author counting GTF, complete gene-filtering code and uncentered per-run TPMs are not supplied. Thus annotation version, representative-only versus all-isoform counting, and the exact background eligibility population remain fidelity limitations. The selected-read primary follows publication-specific Methods; the generic repository uses all unique reads. This run did not change methods to improve agreement.','',
 'Source normalized IP/Input counts already encode the centered enrichment: their log2 ratios agree to floating-point precision. They must not be treated as raw counts or uncentered TPMs. All source count entries are positive. Source profile backgrounds have small nonzero medians, showing that the exported shared gene subset need not equal each original centering population. Neither pseudocounts nor the exact historical eligibility rule can be inferred uniquely.','',
 'TPM uses the union CDS feature length per gene across retained placements. The same lengths apply in IP and input, and both length and the TPM library scale cancel after CYTONUC-median centering; an independent count-ratio identity was verified. Same-strand ambiguous overlaps are assigned to every overlapping gene under --nonunique=all. This option is distinct from genome multimapping, which is excluded by NH=1.','',
 f"Annotation audit: {json.dumps(audit['annotation'])}. Topology audit: {json.dumps(audit['topology'])}.",
 f"All {sum(q['records'] for q in qcs.values()):,} BAM records passed full length, NH=1, mapped, single-end and primary-alignment checks. Input hashes were checked before and after execution. Counts, TPMs, all-gene finite/nonfinite enrichments, background membership and exclusion counts are retained in supporting TSVs.",
 '', '## Profile discrepancy diagnostics','']
 for p in profiles:lines.append(f"- {p['profile']}: {p['systematic_discrepancy']}; residual 5th–95th percentiles {p['residual_p05']:.6f} to {p['residual_p95']:.6f}; sign concordance {p['sign_concordance_nonzero']:.6f}.")
 lines+=['','## Coverage and centering','', '```json',json.dumps(centers,indent=2),'```','',
 '## Reproducibility','',f'Implementation: `{Path(__file__).name}`. Exact run command is recorded in RUN.log; per-BAM HTSeq commands and diagnostic counters are retained. SHA256 input and output manifests accompany the results. Initial HTSeq installation required local build dependencies because the bundled Python lacked a compatible binary wheel; system environments were not modified.','',
 'Publication: https://www.nature.com/articles/s41594-025-01691-6 (Methods: Rfoot-seq data processing, Gene structure annotation, Gene-level expression and enrichment analyses). Generic workflow: https://github.com/zhejilab/RNaseFootprinting (retained README step 9).',
 '', 'Stop here for independent review. Validation 2 and Validation 3 are not authorized in this task.','', '```text',terminal.rstrip(),'```']
 (OUT/(PREFIX+'SUMMARY.md')).write_text('\n'.join(lines)+'\n')

if __name__=='__main__':main()
