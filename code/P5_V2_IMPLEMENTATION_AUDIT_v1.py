from pathlib import Path
import os
from collections import defaultdict
import hashlib
import gzip
import re

import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl import load_workbook


# ============================================================
# PATHS / LOCKS
# ============================================================

SOURCE_PARENT = Path(os.environ["P5_DATA_PARENT"])

ROOT = (
    SOURCE_PARENT
    / "paper5_evolution_spatial_bridge"
    / "05_GSE297497_MPT_BRIDGE"
)

OUTCOME = Path(
    "P5_V2_TMD_OUTCOME_TABLE.tsv"
)

REVEAL_CODE = Path(
    "P5_V2_BUILD_AND_REVEAL.py"
)

SPEC = (
    ROOT
    / "07_PREREVEAL_FREEZE"
    / "P5_V2_PREREVEAL_MPT_MODEL_SPEC.yaml"
)

REFGENE = (
    ROOT
    / "09_FORMAL_REFERENCE"
    / "P5_V2_CURRENT_UCSC_hg38_refGene.txt.gz"
)

RUNROOT = (
    ROOT
    / "12_FORMAL_MPT_FULL"
    / "runs"
)

OCCDIR = Path(
    "occupancy_npz"
)

SOURCE_XLSX = (
    ROOT
    / "01_METADATA_SUPP_AUDIT"
    / "PUBLICATION_SOURCE_DATA"
    / "41594_2025_1691_MOESM3_ESM.xlsx"
)

EXPECTED_SPEC_SHA = (
    "181ed0bef6bc81e4750f421c1e0b3f79"
    "a7bfd5fc55432d60664a191d7854b912"
)

EXPECTED_REVEAL_SHA = (
    "2b0baf7bb00f45af3baecdd9b1790ba1"
    "2867e78ff377d44ecd8902b0a6f660d5"
)

PROFILES = {
    "TMCO1_rep1":
        ("SRR33621384", "SRR33621383"),
    "TMCO1_rep2":
        ("SRR33621382", "SRR33621381"),
    "CCDC47_rep1":
        ("SRR33621380", "SRR33621379"),
    "CCDC47_rep2":
        ("SRR33621378", "SRR33621377"),
    "Nicalin_rep1":
        ("SRR33621376", "SRR33621375"),
    "Nicalin_rep2":
        ("SRR33621374", "SRR33621373"),
}

CANONICAL = {
    *(f"chr{i}" for i in range(1, 23)),
    "chrX",
    "chrY",
    "chrM",
}


def sha(path):
    h = hashlib.sha256()

    with open(path, "rb") as f:
        while True:
            b = f.read(1024 * 1024)

            if not b:
                break

            h.update(b)

    return h.hexdigest()


print("=" * 80)
print("P5-V2 IMPLEMENTATION AUDIT")
print("=" * 80)

print()
print("===== 1. FROZEN IDENTITY =====")

spec_sha = sha(SPEC)
reveal_sha = sha(REVEAL_CODE)

print("SPEC_SHA256 =", spec_sha)
print("REVEAL_CODE_SHA256 =", reveal_sha)

assert spec_sha == EXPECTED_SPEC_SHA
assert reveal_sha == EXPECTED_REVEAL_SHA

print("FROZEN_IDENTITY_GATE = PASS")


# ============================================================
# LOAD OUTCOME
# ============================================================

m = pd.read_csv(
    OUTCOME,
    sep="\t",
)

assert len(m) == 5538

print()
print("===== 2. OUTCOME TABLE =====")
print("OUTCOME_ROWS =", len(m))
print(
    "OUTCOME_GENES =",
    m["human_ensembl_gene"].nunique()
)

if "author_tx" not in m:
    m["author_tx"] = (
        m["author_transcript_names"]
        .astype(str)
        .str.strip()
    )

targets = sorted(
    m["author_tx"].unique()
)

assert len(targets) == 1400

print("TARGET_TRANSCRIPTS =", len(targets))


# ============================================================
# WINDOW OFF-BY-ONE AUDIT
# ============================================================

print()
print("===== 3. WINDOW ARITHMETIC =====")

example_start = 200
example_end = 220

# Primary:
# E+30 ... E+90
ps = example_end + 29
pe = example_end + 90

primary_codons = list(
    range(
        ps + 1,
        pe + 1
    )
)

# Negative:
# S-90 ... S-30
ns = example_start - 91
ne = example_start - 30

negative_codons = list(
    range(
        ns + 1,
        ne + 1
    )
)

print(
    "PRIMARY_EXAMPLE_FIRST_CODON =",
    primary_codons[0]
)

print(
    "PRIMARY_EXAMPLE_LAST_CODON =",
    primary_codons[-1]
)

print(
    "PRIMARY_EXAMPLE_N =",
    len(primary_codons)
)

print(
    "EXPECTED_PRIMARY = 250..310"
)

print(
    "NEGATIVE_EXAMPLE_FIRST_CODON =",
    negative_codons[0]
)

print(
    "NEGATIVE_EXAMPLE_LAST_CODON =",
    negative_codons[-1]
)

print(
    "NEGATIVE_EXAMPLE_N =",
    len(negative_codons)
)

print(
    "EXPECTED_NEGATIVE = 110..170"
)

assert primary_codons[0] == 250
assert primary_codons[-1] == 310
assert len(primary_codons) == 61

assert negative_codons[0] == 110
assert negative_codons[-1] == 170
assert len(negative_codons) == 61

print("WINDOW_OFF_BY_ONE_GATE = PASS")


# ============================================================
# LOAD ALL 12 OCCUPANCY CACHES
# ============================================================

print()
print("===== 4. LOAD OCCUPANCY CACHES =====")

all_runs = sorted(
    {
        run
        for pair in PROFILES.values()
        for run in pair
    }
)

occ = {}

means = {}

for run in all_runs:

    p = (
        OCCDIR
        / f"{run}.target_codon_rpm.npz"
    )

    assert p.exists()

    with np.load(
        p,
        allow_pickle=False,
    ) as z:

        occ[run] = {
            k: z[k]
            for k in z.files
        }

    assert set(occ[run]) == set(targets)

    means[run] = {
        tx: float(
            np.mean(arr)
        )
        for tx, arr
        in occ[run].items()
    }

    print(
        run,
        "TARGETS =",
        len(occ[run])
    )

print("OCCUPANCY_CACHE_GATE = PASS")


# ============================================================
# VALID-UNIVERSE + INDEPENDENT OUTCOME RECONSTRUCTION
# ============================================================

print()
print(
    "===== 5. INDEPENDENT OUTCOME RECONSTRUCTION ====="
)

n = len(m)

calc_primary_profiles = {}
calc_negative_profiles = {}

primary_geometry = np.zeros(
    n,
    dtype=bool,
)

negative_geometry = np.zeros(
    n,
    dtype=bool,
)

all_six_A_positive = np.zeros(
    n,
    dtype=bool,
)

for pname, (
    input_run,
    ip_run,
) in PROFILES.items():

    pv = np.full(
        n,
        np.nan,
        dtype=float,
    )

    nv = np.full(
        n,
        np.nan,
        dtype=float,
    )

    for i, row in enumerate(
        m.itertuples(index=False)
    ):

        tx = row.author_tx

        inp = occ[input_run][tx]
        ip = occ[ip_run][tx]

        A = means[input_run][tx]

        start = int(
            row.author_start_1based
        )

        end = int(
            row.author_end_1based
        )

        plen = len(inp)

        if pname == "TMCO1_rep1":

            primary_geometry[i] = (
                end + 90 <= plen
            )

            negative_geometry[i] = (
                start - 90 >= 1
            )

        if (
            A > 0
            and end + 90 <= plen
        ):

            x = np.log2(
                (
                    ip[
                        end + 29:
                        end + 90
                    ]
                    + A
                )
                /
                (
                    inp[
                        end + 29:
                        end + 90
                    ]
                    + A
                )
            )

            assert len(x) == 61

            pv[i] = float(
                np.mean(x)
            )

        if (
            A > 0
            and start - 90 >= 1
        ):

            x = np.log2(
                (
                    ip[
                        start - 91:
                        start - 30
                    ]
                    + A
                )
                /
                (
                    inp[
                        start - 91:
                        start - 30
                    ]
                    + A
                )
            )

            assert len(x) == 61

            nv[i] = float(
                np.mean(x)
            )

    calc_primary_profiles[pname] = pv
    calc_negative_profiles[pname] = nv


for i, row in enumerate(
    m.itertuples(index=False)
):

    tx = row.author_tx

    all_six_A_positive[i] = all(
        means[input_run][tx] > 0
        for input_run, _
        in PROFILES.values()
    )


def compare_vector(
    name,
    stored,
    calc,
):

    stored = np.asarray(
        stored,
        dtype=float,
    )

    calc = np.asarray(
        calc,
        dtype=float,
    )

    nan_mismatch = int(
        np.sum(
            np.isnan(stored)
            != np.isnan(calc)
        )
    )

    finite = (
        np.isfinite(stored)
        & np.isfinite(calc)
    )

    if finite.any():
        maxdiff = float(
            np.max(
                np.abs(
                    stored[finite]
                    - calc[finite]
                )
            )
        )
    else:
        maxdiff = 0.0

    print(
        name,
        "NAN_STATUS_MISMATCH =",
        nan_mismatch,
    )

    print(
        name,
        "MAX_ABS_DIFF =",
        "%.12g" % maxdiff,
    )

    assert nan_mismatch == 0
    assert maxdiff < 1e-6


for pname in PROFILES:

    compare_vector(
        pname + "_PRIMARY",
        m[
            pname
            + "_TMD_score"
        ].to_numpy(),
        calc_primary_profiles[pname],
    )

    compare_vector(
        pname + "_NEGATIVE",
        m[
            pname
            + "_NEG_score"
        ].to_numpy(),
        calc_negative_profiles[pname],
    )


P = np.column_stack(
    [
        calc_primary_profiles[p]
        for p in PROFILES
    ]
)

N = np.column_stack(
    [
        calc_negative_profiles[p]
        for p in PROFILES
    ]
)

primary_calc = np.where(
    np.isfinite(P).sum(axis=1) == 6,
    np.nanmean(P, axis=1),
    np.nan,
)

negative_calc = np.where(
    np.isfinite(N).sum(axis=1) == 6,
    np.nanmean(N, axis=1),
    np.nan,
)

compare_vector(
    "PRIMARY_SIX_PROFILE_MEAN",
    m[
        "primary_MPT_TMD_score"
    ].to_numpy(),
    primary_calc,
)

compare_vector(
    "NEGATIVE_SIX_PROFILE_MEAN",
    m[
        "temporal_negative_score"
    ].to_numpy(),
    negative_calc,
)


expected_primary_valid = (
    primary_geometry
    & all_six_A_positive
)

expected_negative_valid = (
    negative_geometry
    & all_six_A_positive
)

stored_primary_valid = (
    m[
        "primary_MPT_TMD_score"
    ].notna()
    .to_numpy()
)

stored_negative_valid = (
    m[
        "temporal_negative_score"
    ].notna()
    .to_numpy()
)

print()
print(
    "PRIMARY_FULL_GEOMETRY =",
    int(
        primary_geometry.sum()
    )
)

print(
    "NEGATIVE_FULL_GEOMETRY =",
    int(
        negative_geometry.sum()
    )
)

print(
    "ALL_SIX_INPUT_A_POSITIVE =",
    int(
        all_six_A_positive.sum()
    )
)

print(
    "EXPECTED_PRIMARY_VALID =",
    int(
        expected_primary_valid.sum()
    )
)

print(
    "STORED_PRIMARY_VALID =",
    int(
        stored_primary_valid.sum()
    )
)

print(
    "EXPECTED_NEGATIVE_VALID =",
    int(
        expected_negative_valid.sum()
    )
)

print(
    "STORED_NEGATIVE_VALID =",
    int(
        stored_negative_valid.sum()
    )
)

assert np.array_equal(
    expected_primary_valid,
    stored_primary_valid,
)

assert np.array_equal(
    expected_negative_valid,
    stored_negative_valid,
)

assert stored_primary_valid.sum() == 2796
assert stored_negative_valid.sum() == 2455

print(
    "VALID_UNIVERSE_RECONSTRUCTION_GATE = PASS"
)

print(
    "OUTCOME_FORMULA_RECONSTRUCTION_GATE = PASS"
)


# ============================================================
# ADJUSTED COVARIATE PRESENCE
# ============================================================

print()
print("===== 6. ADJUSTED MODEL COVARIATES =====")

adjusted_cols = [
    "score_z",
    "median_window_aa_identity_z",
    "human_window_GC3_z",
    "human_window_CpG_density_z",
    "p5_length_aa_z",
    "TMD_mean_Kyte_Doolittle_hydropathy_z",
    "p5_tmd_index_z",
    "log1p_downstream_nonTMD_segment_length_aa_z",
    "downstream_segment_lumenal_indicator",
]

for c in adjusted_cols:
    assert c in m.columns
    print(c, "= PRESENT")

adjusted_complete = (
    m[
        ["primary_MPT_TMD_score"]
        + adjusted_cols
    ]
    .replace(
        [np.inf, -np.inf],
        np.nan,
    )
    .dropna()
)

print(
    "ADJUSTED_COMPLETE_CASE_TMD =",
    len(adjusted_complete)
)

assert len(adjusted_complete) == 2796

print(
    "ADJUSTED_COVARIATE_GATE = PASS"
)


# ============================================================
# PRIMARY vs NEGATIVE ON IDENTICAL TMD UNIVERSE
# ============================================================

print()
print(
    "===== 7. TEMPORAL CONTROL — IDENTICAL UNIVERSE ====="
)

shared = m[
    m[
        "primary_MPT_TMD_score"
    ].notna()
    &
    m[
        "temporal_negative_score"
    ].notna()
].copy()

print(
    "SHARED_PRIMARY_NEGATIVE_TMD =",
    len(shared)
)

print(
    "SHARED_PRIMARY_NEGATIVE_GENES =",
    shared[
        "human_ensembl_gene"
    ].nunique()
)


def cluster_slope(
    df,
    ycol,
):

    X = sm.add_constant(
        df[["score_z"]].astype(float),
        has_constant="add",
    )

    fit = sm.OLS(
        df[ycol].astype(float),
        X,
    ).fit(
        cov_type="cluster",
        cov_kwds={
            "groups":
                df[
                    "human_ensembl_gene"
                ],
            "use_correction":
                True,
        },
        use_t=True,
    )

    ci = fit.conf_int().loc[
        "score_z"
    ]

    return (
        float(
            fit.params[
                "score_z"
            ]
        ),
        float(
            fit.pvalues[
                "score_z"
            ]
        ),
        float(ci.iloc[0]),
        float(ci.iloc[1]),
    )


bp, pp, lp, hp = cluster_slope(
    shared,
    "primary_MPT_TMD_score",
)

bn, pn, ln, hn = cluster_slope(
    shared,
    "temporal_negative_score",
)

print(
    "SHARED_PRIMARY_BETA =",
    "%.8f" % bp
)

print(
    "SHARED_PRIMARY_P =",
    "%.8g" % pp
)

print(
    "SHARED_PRIMARY_CI =",
    "[%.8f, %.8f]" % (
        lp,
        hp,
    )
)

print(
    "SHARED_NEGATIVE_BETA =",
    "%.8f" % bn
)

print(
    "SHARED_NEGATIVE_P =",
    "%.8g" % pn
)

print(
    "SHARED_NEGATIVE_CI =",
    "[%.8f, %.8f]" % (
        ln,
        hn,
    )
)

print(
    "SHARED_NEGATIVE_ABS_BETA_WEAKER =",
    str(
        abs(bn) < abs(bp)
    ).upper()
)


# ============================================================
# TARGET-TRANSCRIPT GENOMIC OVERLAP STRUCTURE
# ============================================================

print()
print(
    "===== 8. TARGET TRANSCRIPT OVERLAP AUDIT ====="
)

target_set = set(targets)

ref = defaultdict(list)

with gzip.open(
    REFGENE,
    "rt",
) as fh:

    for line in fh:

        f = line.rstrip("\n").split("\t")

        if len(f) < 11:
            continue

        tx = f[1]

        if (
            tx not in target_set
            or f[2] not in CANONICAL
        ):
            continue

        ref[tx].append(f)


for tx in targets:
    assert len(ref[tx]) == 1


events = defaultdict(
    lambda: defaultdict(int)
)

for tx in targets:

    f = ref[tx]

    f = f[0]

    chrom = f[2]
    strand = f[3]

    starts = [
        int(x)
        for x in f[9]
        .rstrip(",")
        .split(",")
        if x
    ]

    ends = [
        int(x)
        for x in f[10]
        .rstrip(",")
        .split(",")
        if x
    ]

    for s, e in zip(
        starts,
        ends,
    ):

        events[
            (chrom, strand)
        ][s] += 1

        events[
            (chrom, strand)
        ][e] -= 1


def regions_at_least(
    ev,
    threshold,
):

    out = []

    cov = 0
    prev = None

    for pos in sorted(ev):

        if (
            prev is not None
            and prev < pos
            and cov >= threshold
        ):
            out.append(
                (prev, pos)
            )

        cov += ev[pos]
        prev = pos

    return out


union_regions = {}
overlap_regions = {}

for key, ev in events.items():

    union_regions[key] = (
        regions_at_least(
            ev,
            1,
        )
    )

    overlap_regions[key] = (
        regions_at_least(
            ev,
            2,
        )
    )


union_bp = sum(
    e - s
    for rr in union_regions.values()
    for s, e in rr
)

overlap_bp = sum(
    e - s
    for rr in overlap_regions.values()
    for s, e in rr
)

print(
    "TARGET_EXON_UNION_BP =",
    union_bp
)

print(
    "TARGET_EXON_MULTITRANSCRIPT_OVERLAP_BP =",
    overlap_bp
)

print(
    "TARGET_EXON_OVERLAP_BP_FRACTION =",
    "%.8f"
    % (
        overlap_bp / union_bp
        if union_bp
        else 0
    )
)


def compile_regions(d):

    out = {}

    for key, rr in d.items():

        if not rr:
            continue

        out[key] = (
            np.array(
                [x[0] for x in rr],
                dtype=np.int64,
            ),
            np.array(
                [x[1] for x in rr],
                dtype=np.int64,
            ),
        )

    return out


union_arrays = compile_regions(
    union_regions
)

overlap_arrays = compile_regions(
    overlap_regions
)


def count_in_regions(
    positions,
    region_pair,
):

    if (
        region_pair is None
        or len(positions) == 0
    ):
        return 0

    starts, ends = region_pair

    idx = np.searchsorted(
        starts,
        positions,
        side="right",
    ) - 1

    valid = idx >= 0

    if not valid.any():
        return 0

    idxv = idx[valid]
    posv = positions[valid]

    return int(
        np.sum(
            posv < ends[idxv]
        )
    )


total_target = 0
total_overlap = 0
total_asite = 0

for run in all_runs:

    p = (
        RUNROOT
        / run
        / "final"
        / f"{run}.asite_plus15.tsv"
    )

    n_all = 0
    n_target = 0
    n_overlap = 0

    for chunk in pd.read_csv(
        p,
        sep="\t",
        usecols=[
            "chrom",
            "strand",
            "asite_0based",
        ],
        chunksize=1_000_000,
    ):

        n_all += len(chunk)

        for (
            chrom,
            strand,
        ), g in chunk.groupby(
            [
                "chrom",
                "strand",
            ],
            sort=False,
        ):

            pos = g[
                "asite_0based"
            ].to_numpy(
                dtype=np.int64,
            )

            key = (
                str(chrom),
                str(strand),
            )

            n_target += count_in_regions(
                pos,
                union_arrays.get(key),
            )

            n_overlap += count_in_regions(
                pos,
                overlap_arrays.get(key),
            )

    frac = (
        n_overlap / n_target
        if n_target
        else 0
    )

    print(
        run,
        "ALL_ASITE =",
        n_all,
        "TARGET_EXONIC =",
        n_target,
        "OVERLAP =",
        n_overlap,
        "OVERLAP/TARGET =",
        "%.8f" % frac,
    )

    total_asite += n_all
    total_target += n_target
    total_overlap += n_overlap


print(
    "ALL_RUNS_ASITE =",
    total_asite
)

print(
    "ALL_RUNS_TARGET_EXONIC_ASITE =",
    total_target
)

print(
    "ALL_RUNS_AMBIGUOUS_OVERLAP_ASITE =",
    total_overlap
)

print(
    "ALL_RUNS_OVERLAP_AMONG_TARGET_FRACTION =",
    "%.8f"
    % (
        total_overlap / total_target
        if total_target
        else 0
    )
)


# ============================================================
# MPT SOURCE-DATA HEADER CANDIDATES FOR NEXT STEP
# ============================================================

print()
print(
    "===== 9. MPT GENE-LEVEL SOURCE-DATA CANDIDATES ====="
)

wb = load_workbook(
    SOURCE_XLSX,
    read_only=True,
    data_only=True,
)

ws = wb["MPT"]

pattern = re.compile(
    r"gene|transcript|mpt|tmco1|ccdc47|nicalin|"
    r"enrich|input|ip|tpm|log",
    re.I,
)

for r_idx, row in enumerate(
    ws.iter_rows(
        min_row=1,
        max_row=8,
        values_only=True,
    ),
    start=1,
):

    for c_idx, value in enumerate(
        row,
        start=1,
    ):

        if (
            value is not None
            and pattern.search(
                str(value)
            )
        ):

            print(
                "ROW",
                r_idx,
                "COL",
                c_idx,
                "=",
                repr(value),
            )


print()
print("=" * 80)

print(
    "IMPLEMENTATION_CORE_AUDIT = PASS"
)

print(
    "PRIMARY_RESULT_NOT_RETUNED = TRUE"
)

print(
    "MPT_OUTCOME_ALREADY_REVEALED = TRUE"
)

print(
    "NEXT = GENE_LEVEL_COMPANION_AFTER_AUDIT_REVIEW"
)

print("=" * 80)
