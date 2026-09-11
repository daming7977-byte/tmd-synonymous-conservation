from pathlib import Path
import os
import hashlib
import numpy as np
import pandas as pd
import statsmodels.api as sm
from scipy.stats import spearmanr
from openpyxl import load_workbook

SOURCE_PARENT = Path(os.environ["P5_DATA_PARENT"])

ROOT = (
    SOURCE_PARENT
    / "paper5_evolution_spatial_bridge"
    / "05_GSE297497_MPT_BRIDGE"
)

OUTCOME = Path(
    "P5_V2_TMD_OUTCOME_TABLE.tsv"
)

SOURCE = (
    ROOT
    / "01_METADATA_SUPP_AUDIT"
    / "PUBLICATION_SOURCE_DATA"
    / "41594_2025_1691_MOESM3_ESM.xlsx"
)

SPEC = (
    ROOT
    / "07_PREREVEAL_FREEZE"
    / "P5_V2_PREREVEAL_MPT_MODEL_SPEC.yaml"
)

EXPECTED_SPEC_SHA = (
    "181ed0bef6bc81e4750f421c1e0b3f79"
    "a7bfd5fc55432d60664a191d7854b912"
)

OUTCOMES = [
    "MPT average enrichment (log2)",
    "TMCO1 average enrichment (log2)",
    "CCDC47 average enrichment (log2)",
    "Nicalin average enrichment (log2)",
]


def sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1024*1024), b""):
            h.update(b)
    return h.hexdigest()


def holm(pvals):
    pvals = np.asarray(pvals, dtype=float)
    m = len(pvals)
    order = np.argsort(pvals)
    out = np.empty(m)
    running = 0.0

    for rank, idx in enumerate(order):
        val = min(
            1.0,
            pvals[idx] * (m-rank)
        )
        val = max(val, running)
        running = val
        out[idx] = val

    return out


def fit_gene_level(df, ycol):

    d = (
        df[
            [
                "gene_mean_evo_z",
                ycol,
            ]
        ]
        .replace(
            [np.inf, -np.inf],
            np.nan
        )
        .dropna()
        .copy()
    )

    X = sm.add_constant(
        d[["gene_mean_evo_z"]],
        has_constant="add"
    )

    fit = sm.OLS(
        d[ycol].astype(float),
        X.astype(float)
    ).fit(
        cov_type="HC3"
    )

    ci = fit.conf_int().loc[
        "gene_mean_evo_z"
    ]

    rho, sp = spearmanr(
        d["gene_mean_evo_z"],
        d[ycol],
        nan_policy="omit"
    )

    return {
        "N": len(d),
        "beta": float(
            fit.params["gene_mean_evo_z"]
        ),
        "se": float(
            fit.bse["gene_mean_evo_z"]
        ),
        "ci_low": float(ci.iloc[0]),
        "ci_high": float(ci.iloc[1]),
        "p": float(
            fit.pvalues["gene_mean_evo_z"]
        ),
        "spearman_rho": float(rho),
        "spearman_p": float(sp),
    }


print("="*80)
print("P5-V2 GENE-LEVEL COMPANION")
print("="*80)

spec_sha = sha(SPEC)

print(
    "PREREVEAL_SPEC_SHA256 =",
    spec_sha
)

assert spec_sha == EXPECTED_SPEC_SHA


# ============================================================
# TMD evolutionary side
# ============================================================

m = pd.read_csv(
    OUTCOME,
    sep="\t"
)

assert len(m) == 5538
assert (
    m["human_ensembl_gene"]
    .nunique()
    == 1400
)

# Recover exact author gene name using
# transcript -> Gene mapping in DeepTMHMM.

target_tx = set(
    m["author_tx"].astype(str)
)

wb = load_workbook(
    SOURCE,
    read_only=True,
    data_only=True
)

ws = wb["DeepTMHMM topology"]

tx_to_gene = {}

for row in ws.iter_rows(
    min_row=2,
    max_col=2,
    values_only=True
):

    gene, tx = row

    if tx is None:
        continue

    tx = str(tx).strip()

    if tx not in target_tx:
        continue

    gene = str(gene).strip()

    if (
        tx in tx_to_gene
        and tx_to_gene[tx] != gene
    ):
        raise RuntimeError(
            "TRANSCRIPT_GENE_CONFLICT: "
            + tx
        )

    tx_to_gene[tx] = gene


assert set(tx_to_gene) == target_tx

m["source_gene"] = (
    m["author_tx"]
    .astype(str)
    .map(tx_to_gene)
)

# Use ALL prelocked TMDs per gene,
# not only MPT-valid TMDs.

gene = (
    m.groupby(
        [
            "human_ensembl_gene",
            "source_gene"
        ],
        as_index=False
    )
    .agg(
        gene_mean_score_z=(
            "score_z",
            "mean"
        ),
        n_prelocked_TMD=(
            "p5_tmd_index",
            "size"
        )
    )
)

assert len(gene) == 1400

# Re-standardize gene means across all 1400 genes,
# before looking at Source Data MPT values.

mu = gene[
    "gene_mean_score_z"
].mean()

sd = gene[
    "gene_mean_score_z"
].std(ddof=0)

gene["gene_mean_evo_z"] = (
    gene["gene_mean_score_z"]
    - mu
) / sd

print(
    "GENE_EVOLUTION_UNIVERSE =",
    len(gene)
)

print(
    "GENE_EVOLUTION_AGGREGATION = "
    "MEAN_OF_ALL_PRELOCKED_TMD_SCORE_Z"
)


# ============================================================
# Publication MPT gene-level source data
# ============================================================

src = pd.read_excel(
    SOURCE,
    sheet_name="MPT"
)

print(
    "SOURCE_MPT_ROWS =",
    len(src)
)

assert "Gene" in src.columns

for c in OUTCOMES:
    assert c in src.columns


dup = int(
    src["Gene"]
    .astype(str)
    .duplicated()
    .sum()
)

print(
    "SOURCE_DUPLICATE_GENE_ROWS =",
    dup
)

if dup != 0:
    raise RuntimeError(
        "SOURCE_MPT_GENE_NOT_UNIQUE"
    )

src = src[
    ["Gene"] + OUTCOMES
].copy()

src["Gene"] = (
    src["Gene"]
    .astype(str)
    .str.strip()
)

for c in OUTCOMES:
    src[c] = pd.to_numeric(
        src[c],
        errors="coerce"
    )


d = gene.merge(
    src,
    left_on="source_gene",
    right_on="Gene",
    how="left",
    validate="one_to_one"
)

print(
    "TARGET_GENES_MATCHED_TO_MPT_SOURCE =",
    int(d["Gene"].notna().sum())
)

print(
    "TARGET_GENES_MISSING_FROM_MPT_SOURCE =",
    int(d["Gene"].isna().sum())
)

print(
    "TARGET_GENES_WITH_VALID_MPT_AVERAGE =",
    int(
        d[
            "MPT average enrichment (log2)"
        ]
        .notna()
        .sum()
    )
)


# ============================================================
# Primary companion
# ============================================================

primary = fit_gene_level(
    d,
    "MPT average enrichment (log2)"
)

factor_results = {}

for factor in [
    "TMCO1",
    "CCDC47",
    "Nicalin"
]:

    factor_results[factor] = (
        fit_gene_level(
            d,
            factor
            + " average enrichment (log2)"
        )
    )


ph = holm(
    [
        factor_results[f]["p"]
        for f in [
            "TMCO1",
            "CCDC47",
            "Nicalin"
        ]
    ]
)

for f, p in zip(
    [
        "TMCO1",
        "CCDC47",
        "Nicalin"
    ],
    ph
):
    factor_results[f]["holm_p"] = float(p)


# ============================================================
# Interpretation
# ============================================================

tmd_primary_beta = -0.0212987

same_direction = (
    np.sign(primary["beta"])
    ==
    np.sign(tmd_primary_beta)
)

global_gene_preference = (
    primary["p"] < 0.05
    and same_direction
)


def fmt(r):
    return (
        "N={N} beta={beta:.6g} "
        "SE={se:.6g} "
        "95%CI=[{ci_low:.6g},{ci_high:.6g}] "
        "p={p:.6g} "
        "Spearman_rho={spearman_rho:.6g} "
        "Spearman_p={spearman_p:.6g}"
    ).format(**r)


print()
print("="*80)
print("GENE-LEVEL COMPANION REVEAL")
print("="*80)

print(
    "PUBLICATION_MPT_AVERAGE:",
    fmt(primary)
)

print()
print("FACTOR-SPECIFIC:")

for f in [
    "TMCO1",
    "CCDC47",
    "Nicalin"
]:

    r = factor_results[f]

    print(
        f + ":",
        fmt(r),
        "Holm_p=%.6g"
        % r["holm_p"]
    )


print()

print(
    "GENE_LEVEL_DIRECTION_SAME_AS_TMD_PRIMARY =",
    str(same_direction).upper()
)

print(
    "SIGNIFICANT_SAME_DIRECTION_GLOBAL_GENE_PREFERENCE =",
    str(global_gene_preference).upper()
)


if global_gene_preference:

    interpretation = (
        "GENE_LEVEL_PREFERENCE_PRESENT__"
        "WEAKENS_LOCAL_TMD_SPECIFIC_INTERPRETATION"
    )

elif primary["p"] >= 0.05:

    interpretation = (
        "NO_SIGNIFICANT_GENE_LEVEL_MPT_BRIDGE__"
        "GLOBAL_GENE_PREFERENCE_NOT_SUPPORTED"
    )

else:

    interpretation = (
        "GENE_LEVEL_ASSOCIATION_DIFFERENT_DIRECTION"
    )


print(
    "GENE_LEVEL_COMPANION_INTERPRETATION =",
    interpretation
)

print()
print(
    "FORMAL_P5_V2_PRIMARY_ADJUDICATION_REMAINS = "
    "INTERMEDIATE_BRIDGE_RESULT"
)

print(
    "PRIMARY_RESULT_RETUNED = FALSE"
)

print(
    "GENE_LEVEL_COMPANION_ROLE = SECONDARY"
)


# ============================================================
# OUTPUTS
# ============================================================

d.to_csv(
    "P5_V2_GENE_LEVEL_COMPANION_DATA.tsv",
    sep="\t",
    index=False
)

rows = [
    {
        "analysis":
            "MPT_average",
        **primary
    }
]

for f in [
    "TMCO1",
    "CCDC47",
    "Nicalin"
]:

    rows.append(
        {
            "analysis":
                f + "_average",
            **factor_results[f]
        }
    )

pd.DataFrame(
    rows
).to_csv(
    "P5_V2_GENE_LEVEL_COMPANION_RESULTS.tsv",
    sep="\t",
    index=False
)

print()
print(
    "GENE_LEVEL_COMPANION_OUTPUTS_WRITTEN = TRUE"
)
